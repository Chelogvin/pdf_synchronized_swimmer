"""
PDF Переводчик — с автоматической сменой темы (изумрудная палитра)
Требует: PyMuPDF, deep-translator, Pillow, python-docx, requests, sv-ttk
"""

import sys
import subprocess

def install_if_missing(package, import_name=None):
    import_name = import_name or package
    try:
        __import__(import_name)
    except ImportError:
        print(f"Устанавливаю {package}…")
        subprocess.check_call([sys.executable, "-m", "pip", "install", package, "-q"])

install_if_missing("PyMuPDF", "fitz")
install_if_missing("deep-translator", "deep_translator")
install_if_missing("Pillow", "PIL")
install_if_missing("python-docx", "docx")
install_if_missing("requests")
install_if_missing("sv-ttk", "sv_ttk")

import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext
import fitz
from PIL import Image, ImageTk
from deep_translator import GoogleTranslator
import docx as python_docx
import re
import threading
import io
import requests
import json
import sv_ttk
from datetime import datetime

# ─────────────────────────────────────────────────────────────────────────────
# Константы
LANGUAGES = {
    "Русский": "ru",
    "Английский": "en",
    "Немецкий": "de",
    "Французский": "fr",
    "Испанский": "es",
    "Итальянский": "it",
    "Китайский": "zh-CN",
    "Японский": "ja",
}

DEEPSEEK_URL = "https://api.deepseek.com/v1/chat/completions"
DEEPSEEK_MODEL = "deepseek-chat"

# Изумрудная палитра
EMERALD = {
    "light": {
        "bg": "#f0faf5",
        "toolbar": "#e0f0e8",
        "btn": "#1a6e4a",
        "btn_fg": "#ffffff",
        "canvas_bg": "#d4e6dd",
        "text_bg": "#ffffff",
        "text_fg": "#1a2a24",
        "accent": "#2ecc71",
        "nav_bg": "#e0f0e8",
        "border": "#1a6e4a",
        "status": "#555555",
    },
    "dark": {
        "bg": "#0a120f",
        "toolbar": "#0d2a1f",
        "btn": "#1a6e4a",
        "btn_fg": "#ffffff",
        "canvas_bg": "#0d1f18",
        "text_bg": "#1a1f1c",
        "text_fg": "#e0f2e9",
        "accent": "#2ecc71",
        "nav_bg": "#0d2a1f",
        "border": "#2ecc71",
        "status": "#88998f",
    }
}

SENT_COLORS_LIGHT = ["#e8f5e9", "#e0f2e9", "#e3f2fd", "#fce4ec", "#e0f7fa"]
SENT_COLORS_DARK = ["#1a2a24", "#1d3b30", "#143028", "#1e3028", "#1a3528"]


class PDFTranslatorApp:
    def __init__(self, root: tk.Tk):
        self.root = root
        self.root.title("PDF Переводчик")
        self.root.geometry("1400x860")
        self.root.minsize(900, 600)

        # Определяем тему по времени (20:00 – 6:00 — тёмная)
        current_hour = datetime.now().hour
        self.current_theme = "dark" if current_hour >= 20 or current_hour < 6 else "light"

        # Применяем sv-ttk тему как основу
        if self.current_theme == "dark":
            sv_ttk.set_theme("dark")
        else:
            sv_ttk.set_theme("light")

        # Переопределяем цвета вручную (изумрудные)
        self._apply_custom_colors()

        # PDF
        self.pdf_doc = None
        self.pdf_path = ""
        self.current_page = 0
        self.zoom = 1.5
        self.page_image = None
        self.page_photo = None
        self.image_item = None

        # Выделение
        self.sel_start = None
        self.sel_rect_id = None
        self.selection_rect = None
        self.highlight_rect_id = None

        # Прокрутка
        self.sync_scroll_var = tk.BooleanVar(value=True)
        self._scrolling = False

        # Перевод
        self.deepseek_key = tk.StringVar(value="")
        self.target_lang_name = tk.StringVar(value="Русский")

        # Обрезка
        self.crop_top = tk.IntVar(value=0)
        self.crop_bottom = tk.IntVar(value=0)
        self._crop_top_line_id = None
        self._crop_bottom_line_id = None

        # Всплывающее окно
        self.popup_win = None

        # Таймер и флаг
        self.translate_timer = None
        self.is_translating = False

        self._build_ui()
        self._bind_hotkeys()

        # начальный размер шрифта
        self.trans_font_size = 13

    # ════════════════════════════════════════════════════════════════════════
    # ЦВЕТА
    def _apply_custom_colors(self):
        colors = EMERALD[self.current_theme]
        self.root.configure(bg=colors["bg"])
        self.colors = colors

    def _get_sentence_colors(self):
        return SENT_COLORS_DARK if self.current_theme == "dark" else SENT_COLORS_LIGHT

    # ════════════════════════════════════════════════════════════════════════
    # UI
    def _build_ui(self):
        colors = self.colors

        # Верхняя панель (ttk)
        toolbar = ttk.Frame(self.root)
        toolbar.pack(side=tk.TOP, fill=tk.X, padx=5, pady=5)

        ttk.Button(toolbar, text="📂 Открыть PDF", command=self.open_pdf).pack(side=tk.LEFT, padx=2)
        ttk.Button(toolbar, text="💾 Сохранить DOCX", command=self.save_docx).pack(side=tk.LEFT, padx=2)

        ttk.Label(toolbar, text="Язык:").pack(side=tk.LEFT, padx=(10, 2))
        lang_combo = ttk.Combobox(toolbar, textvariable=self.target_lang_name,
                                  values=list(LANGUAGES.keys()), state="readonly", width=12)
        lang_combo.pack(side=tk.LEFT, padx=2)

        ttk.Checkbutton(toolbar, text="Синхр. прокрутка", variable=self.sync_scroll_var).pack(side=tk.LEFT, padx=10)

        ttk.Label(toolbar, text="Масштаб:").pack(side=tk.LEFT, padx=(10, 2))
        ttk.Button(toolbar, text="−", command=self.zoom_out, width=3).pack(side=tk.LEFT)
        self.zoom_label = ttk.Label(toolbar, text="150%", width=5)
        self.zoom_label.pack(side=tk.LEFT)
        ttk.Button(toolbar, text="+", command=self.zoom_in, width=3).pack(side=tk.LEFT)
        # Разделитель
        ttk.Separator(toolbar, orient=tk.VERTICAL).pack(side=tk.LEFT, padx=10, fill=tk.Y)
        
        ttk.Label(toolbar, text="Текст:").pack(side=tk.LEFT, padx=(5, 2))
        ttk.Button(toolbar, text="A−", command=self.zoom_text_out, width=3).pack(side=tk.LEFT)
        ttk.Button(toolbar, text="A+", command=self.zoom_text_in, width=3).pack(side=tk.LEFT)	

        ttk.Label(toolbar, text="DeepSeek API:").pack(side=tk.LEFT, padx=(15, 2))
        ttk.Entry(toolbar, textvariable=self.deepseek_key, width=28, show="*").pack(side=tk.LEFT)

        # Панель обрезки
        crop_frame = ttk.Frame(self.root)
        crop_frame.pack(side=tk.TOP, fill=tk.X, padx=5, pady=2)

        ttk.Label(crop_frame, text="✂ Обрезка (не переводить):", font=("Segoe UI", 10, "bold")).pack(side=tk.LEFT, padx=5)
        ttk.Label(crop_frame, text="Сверху (px):").pack(side=tk.LEFT, padx=(10, 2))
        top_spin = ttk.Spinbox(crop_frame, from_=0, to=5000, width=6, textvariable=self.crop_top,
                               command=self._update_crop_lines)
        top_spin.pack(side=tk.LEFT, padx=2)

        ttk.Label(crop_frame, text="Снизу (px):").pack(side=tk.LEFT, padx=(10, 2))
        bot_spin = ttk.Spinbox(crop_frame, from_=0, to=5000, width=6, textvariable=self.crop_bottom,
                               command=self._update_crop_lines)
        bot_spin.pack(side=tk.LEFT, padx=2)

        ttk.Button(crop_frame, text="Сбросить", command=self._reset_crop).pack(side=tk.LEFT, padx=10)

        # Панель выделения
        sel_frame = ttk.Frame(self.root)
        sel_frame.pack(side=tk.TOP, fill=tk.X, padx=5, pady=2)

        ttk.Button(sel_frame, text="🔤 Перевести выделенное", command=self.translate_selection).pack(side=tk.LEFT, padx=2)
        ttk.Button(sel_frame, text="🤖 Объяснить ИИ", command=self.explain_ai).pack(side=tk.LEFT, padx=2)
        self.sel_status = ttk.Label(sel_frame, text="Выделите область на PDF")
        self.sel_status.pack(side=tk.LEFT, padx=15)

        # Основная область
        self.paned = ttk.PanedWindow(self.root, orient=tk.HORIZONTAL)
        self.paned.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)

        # Левая панель (PDF)
        left_frame = ttk.Frame(self.paned)
        self.paned.add(left_frame, weight=1)

        self.pdf_canvas = tk.Canvas(left_frame, bg=colors["canvas_bg"], highlightthickness=0, cursor="crosshair")
        self.pdf_canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        pdf_scroll = ttk.Scrollbar(left_frame, orient=tk.VERTICAL, command=self._on_pdf_scroll_y)
        pdf_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.pdf_canvas.configure(yscrollcommand=pdf_scroll.set)

        # Правая панель (текст)
        right_frame = ttk.Frame(self.paned)
        self.paned.add(right_frame, weight=1)

        text_frame = ttk.Frame(right_frame)
        text_frame.pack(fill=tk.BOTH, expand=True)

        self.trans_text = tk.Text(text_frame, wrap=tk.WORD, bg=colors["text_bg"], fg=colors["text_fg"],
                                  font=("Segoe UI", 13), padx=20, pady=20,
                                  insertbackground=colors["text_fg"],
                                  selectbackground=colors["accent"],
                                  relief=tk.FLAT, state=tk.DISABLED)
        self.trans_text.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        text_scroll = ttk.Scrollbar(text_frame, orient=tk.VERTICAL, command=self._on_trans_scroll_y)
        text_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.trans_text.configure(yscrollcommand=text_scroll.set)

        # Нижняя навигация
        nav = ttk.Frame(self.root)
        nav.pack(side=tk.BOTTOM, fill=tk.X, padx=5, pady=5)

        ttk.Button(nav, text="◀ Пред.", command=self.prev_page).pack(side=tk.LEFT, padx=2)
        self.page_label = ttk.Label(nav, text="Страница — / —")
        self.page_label.pack(side=tk.LEFT, padx=10)

        ttk.Label(nav, text="Страница:").pack(side=tk.LEFT, padx=(10, 2))
        self.page_entry = ttk.Entry(nav, width=5)
        self.page_entry.pack(side=tk.LEFT, padx=2)
        self.page_entry.bind("<Return>", self.go_to_page)
        ttk.Button(nav, text="Перейти", command=self.go_to_page).pack(side=tk.LEFT, padx=2)

        ttk.Button(nav, text="След. ▶", command=self.next_page).pack(side=tk.LEFT, padx=2)

        self.status_bar = ttk.Label(nav, text="Откройте PDF-файл", foreground=colors["status"])
        self.status_bar.pack(side=tk.LEFT, padx=20, fill=tk.X, expand=True)

        # Настройка цветов для текста
        self.trans_text.tag_configure("status", foreground=colors["status"], font=("Segoe UI", 11, "italic"))

        # События мыши
        self.pdf_canvas.bind("<ButtonPress-1>", self._sel_start)
        self.pdf_canvas.bind("<B1-Motion>", self._sel_drag)
        self.pdf_canvas.bind("<ButtonRelease-1>", self._sel_end)
        self.pdf_canvas.bind("<ButtonRelease-3>", self._on_right_click)
        self.pdf_canvas.bind("<MouseWheel>", self._on_mousewheel_pdf)
        self.pdf_canvas.bind("<Button-4>", self._on_mousewheel_pdf)
        self.pdf_canvas.bind("<Button-5>", self._on_mousewheel_pdf)

    def _bind_hotkeys(self):
        self.root.bind("<Left>", lambda e: self.prev_page())
        self.root.bind("<Right>", lambda e: self.next_page())
        self.root.bind("<Prior>", lambda e: self.prev_page())
        self.root.bind("<Next>", lambda e: self.next_page())
        self.root.bind("<Control-plus>", lambda e: self.zoom_in())
        self.root.bind("<Control-minus>", lambda e: self.zoom_out())
        self.root.bind("<Control-plus>", lambda e: self.zoom_text_in())
        self.root.bind("<Control-minus>", lambda e: self.zoom_text_out())

    # ════════════════════════════════════════════════════════════════════════
    # НАВИГАЦИЯ
    # ════════════════════════════════════════════════════════════════════════

    def go_to_page(self, event=None):
        if not self.pdf_doc:
            return
        try:
            page_num = int(self.page_entry.get()) - 1
            if 0 <= page_num < len(self.pdf_doc):
                self.current_page = page_num
                self._clear_translation()
                self.render_page()
                self.page_entry.delete(0, tk.END)
                self.page_entry.insert(0, str(page_num + 1))
            else:
                messagebox.showwarning("Вне диапазона", f"Страница от 1 до {len(self.pdf_doc)}")
        except ValueError:
            messagebox.showwarning("Ошибка", "Введите число")

    def prev_page(self):
        if self.pdf_doc and self.current_page > 0:
            self.current_page -= 1
            self._clear_translation()
            self.render_page()

    def next_page(self):
        if self.pdf_doc and self.current_page < len(self.pdf_doc) - 1:
            self.current_page += 1
            self._clear_translation()
            self.render_page()

    # ════════════════════════════════════════════════════════════════════════
    # ОТОБРАЖЕНИЕ PDF
    # ════════════════════════════════════════════════════════════════════════

    def open_pdf(self):
        path = filedialog.askopenfilename(filetypes=[("PDF файлы", "*.pdf")])
        if not path:
            return
        try:
            self.pdf_doc = fitz.open(path)
            self.pdf_path = path
            self.current_page = 0
            self._clear_translation()
            self.render_page()
            self.status_bar.config(text=f"Открыт: {path}")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось открыть PDF:\n{e}")

    def render_page(self):
        if not self.pdf_doc:
            return
        page = self.pdf_doc[self.current_page]
        mat = fitz.Matrix(self.zoom, self.zoom)
        pix = page.get_pixmap(matrix=mat, alpha=False)
        pil_img = Image.open(io.BytesIO(pix.tobytes("ppm")))
        self.page_image = pil_img
        self.page_photo = ImageTk.PhotoImage(pil_img)
        w, h = pil_img.size
        self.pdf_canvas.config(scrollregion=(0, 0, w, h))
        if self.image_item:
            self.pdf_canvas.delete(self.image_item)
        self.image_item = self.pdf_canvas.create_image(0, 0, anchor=tk.NW, image=self.page_photo)
        self._clear_selection()
        self._clear_highlight()
        self._crop_top_line_id = None
        self._crop_bottom_line_id = None
        self._update_crop_lines()
        n = len(self.pdf_doc)
        self.page_label.config(text=f"Страница {self.current_page + 1} / {n}")
        if hasattr(self, 'page_entry'):
            self.page_entry.delete(0, tk.END)
            self.page_entry.insert(0, str(self.current_page + 1))
        self.zoom_label.config(text=f"{int(self.zoom * 100)}%")

        # Отложенный перевод
        if self.translate_timer:
            self.root.after_cancel(self.translate_timer)
        self.translate_timer = self.root.after(300, self.translate_current_page)

    def zoom_in(self):
        if self.zoom < 4.0:
            self.zoom = round(self.zoom + 0.25, 2)
            self.render_page()

    def zoom_out(self):
        if self.zoom > 0.5:
            self.zoom = round(self.zoom - 0.25, 2)
            self.render_page()

    # ════════════════════════════════════════════════════════════════════════
    # ОБРЕЗКА
    # ════════════════════════════════════════════════════════════════════════

    def _update_crop_lines(self):
        if self._crop_top_line_id:
            self.pdf_canvas.delete(self._crop_top_line_id)
        if self._crop_bottom_line_id:
            self.pdf_canvas.delete(self._crop_bottom_line_id)
        if not self.page_image:
            return
        w, h = self.page_image.size
        top_px = self.crop_top.get()
        bottom_px = self.crop_bottom.get()
        if top_px > 0:
            self._crop_top_line_id = self.pdf_canvas.create_line(0, top_px, w, top_px, fill="#e53935", width=2, dash=(8, 4))
        bottom_y = h - bottom_px
        if bottom_px > 0 and bottom_y > 0:
            self._crop_bottom_line_id = self.pdf_canvas.create_line(0, bottom_y, w, bottom_y, fill="#e53935", width=2, dash=(8, 4))

    def _reset_crop(self):
        self.crop_top.set(0)
        self.crop_bottom.set(0)
        self._update_crop_lines()

    def _get_crop_rect_pdf(self):
        if not self.pdf_doc or not self.page_image:
            return None
        top_px = self.crop_top.get()
        bottom_px = self.crop_bottom.get()
        if top_px == 0 and bottom_px == 0:
            return None
        _, img_h = self.page_image.size
        y0 = top_px / self.zoom
        y1 = (img_h - bottom_px) / self.zoom
        page = self.pdf_doc[self.current_page]
        return fitz.Rect(0, y0, page.rect.width, max(y1, y0 + 1))

    # ════════════════════════════════════════════════════════════════════════
    # ВЫДЕЛЕНИЕ
    # ════════════════════════════════════════════════════════════════════════

    def _canvas_to_img(self, cx, cy):
        return int(self.pdf_canvas.canvasx(cx)), int(self.pdf_canvas.canvasy(cy))

    def _sel_start(self, event):
        self._clear_selection()
        self.sel_start = self._canvas_to_img(event.x, event.y)

    def _sel_drag(self, event):
        if not self.sel_start:
            return
        cur = self._canvas_to_img(event.x, event.y)
        if self.sel_rect_id:
            self.pdf_canvas.delete(self.sel_rect_id)
        x0, y0 = self.sel_start
        x1, y1 = cur
        self.sel_rect_id = self.pdf_canvas.create_rectangle(x0, y0, x1, y1, outline="#e6a817", width=2, fill="#FFD700", stipple="gray25")

    def _sel_end(self, event):
        if not self.sel_start:
            return
        x0, y0 = self.sel_start
        x1, y1 = self._canvas_to_img(event.x, event.y)
        self.selection_rect = (min(x0, x1), min(y0, y1), max(x0, x1), max(y0, y1))
        w = self.selection_rect[2] - self.selection_rect[0]
        h = self.selection_rect[3] - self.selection_rect[1]
        if w < 5 or h < 5:
            self._clear_selection()
            return
        self.sel_status.config(text=f"Выделено: {w}×{h} пикс.")
        self.sel_start = None

    def _clear_selection(self):
        if self.sel_rect_id:
            self.pdf_canvas.delete(self.sel_rect_id)
            self.sel_rect_id = None
        self.selection_rect = None
        self.sel_start = None
        self.sel_status.config(text="Выделите область на PDF")

    def _get_selection_text(self):
        if not self.pdf_doc or not self.selection_rect:
            return ""
        x0, y0, x1, y1 = self.selection_rect
        rect = fitz.Rect(x0 / self.zoom, y0 / self.zoom, x1 / self.zoom, y1 / self.zoom)
        return self.pdf_doc[self.current_page].get_text("text", clip=rect).strip()

    # ════════════════════════════════════════════════════════════════════════
    # ТЕКСТ И АБЗАЦЫ
    # ════════════════════════════════════════════════════════════════════════

    def _get_page_text(self):
        if not self.pdf_doc:
            return ""
        clip_rect = self._get_crop_rect_pdf()
        page = self.pdf_doc[self.current_page]
        if clip_rect:
            raw_text = page.get_text("text", clip=clip_rect)
        else:
            raw_text = page.get_text("text")
        if not raw_text.strip():
            return ""
        return self._restore_paragraphs_by_indent(raw_text, page, clip_rect)

    def _restore_paragraphs_by_indent(self, text, page, clip_rect=None):
        if not page:
            return text
        blocks = page.get_text("dict")["blocks"]
        lines_with_info = []
        for block in blocks:
            if block["type"] == 0:
                for line in block["lines"]:
                    x0 = line["bbox"][0]
                    y0 = line["bbox"][1]
                    if clip_rect:
                        if y0 < clip_rect.y0 or y0 > clip_rect.y1:
                            continue
                        x0 = x0 - clip_rect.x0
                    text_line = " ".join([span["text"] for span in line["spans"]])
                    if text_line.strip():
                        lines_with_info.append((x0, text_line))
        if len(lines_with_info) < 2:
            return text
        indents = sorted([x0 for x0, _ in lines_with_info])
        normal_indent = indents[len(indents) // 2]
        paragraph_threshold = normal_indent + 5
        result = []
        for i, (x0, line_text) in enumerate(lines_with_info):
            if i > 0 and x0 > paragraph_threshold:
                result.append("")
            result.append(line_text)
        return '\n'.join(result)

    # ════════════════════════════════════════════════════════════════════════
    # ПЕРЕВОД
    # ════════════════════════════════════════════════════════════════════════

    def _target_lang_code(self):
        return LANGUAGES.get(self.target_lang_name.get(), "ru")

    def _do_translate(self, text, src="auto"):
        if not text or not text.strip():
            return ""
        try:
            paragraphs = re.split(r'\n\s*\n', text.strip())
            translated = []
            translator = GoogleTranslator(source=src, target=self._target_lang_code())
            for para in paragraphs:
                if not para.strip():
                    translated.append("")
                    continue
                cleaned = ' '.join(para.split())
                if cleaned:
                    try:
                        trans = translator.translate(cleaned)
                        translated.append(trans if trans else "")
                    except Exception as e:
                        translated.append("")
            safe = [p for p in translated if p is not None]
            return '\n\n'.join(safe)
        except Exception as e:
            return f"[Ошибка: {e}]"

    def translate_current_page(self):
        if not self.pdf_doc:
            return
        if self.is_translating:
            return
        self._set_status("Перевожу…")
        threading.Thread(target=self._translate_fast, daemon=True).start()

    def _translate_fast(self):
        if self.is_translating:
            return
        self.is_translating = True
        try:
            text = self._get_page_text()
            if not text.strip():
                self.root.after(0, lambda: self._show_translation("Текст не найден"))
                return
            result = self._do_translate(text)
            self.root.after(0, lambda: self._show_translation(result))
            self.root.after(0, lambda: self._set_status("Перевод готов"))
        except Exception as e:
            print(f"Ошибка: {e}")
        finally:
            self.is_translating = False

    def _show_translation(self, text):
        self.trans_text.config(state=tk.NORMAL)
        self.trans_text.delete("1.0", tk.END)
        self.trans_text.insert(tk.END, text)
        self.trans_text.config(state=tk.DISABLED)

    def _clear_translation(self):
        self.trans_text.config(state=tk.NORMAL)
        self.trans_text.delete("1.0", tk.END)
        self.trans_text.config(state=tk.DISABLED)

    def _highlight_in_pdf(self, coords):
        self._clear_highlight()
        if not coords:
            return
        x0, y0, x1, y1 = [c * self.zoom for c in coords]
        pad = 30
        rx0, ry0 = max(0, x0 - pad), max(0, y0 - 4)
        rx1, ry1 = x1 + pad * 4, y1 + 4
        self.highlight_rect_id = self.pdf_canvas.create_rectangle(rx0, ry0, rx1, ry1, outline="#e6a817", width=3, fill="#FFD700", stipple="gray12")
        _, h = self.page_image.size if self.page_image else (1, 1)
        if h > 0:
            frac = (ry0 - self.pdf_canvas.winfo_height() / 4) / h
            self.pdf_canvas.yview_moveto(max(0, frac))

    def _clear_highlight(self):
        if self.highlight_rect_id:
            self.pdf_canvas.delete(self.highlight_rect_id)
            self.highlight_rect_id = None

    def zoom_text_in(self):
        """Увеличить шрифт в правой панели"""
        if self.trans_font_size < 30:
            self.trans_font_size += 1
            self._apply_text_font()

    def zoom_text_out(self):
        """Уменьшить шрифт в правой панели"""
        if self.trans_font_size > 8:
            self.trans_font_size -= 1
            self._apply_text_font()

    def _apply_text_font(self):
        """Применить текущий размер шрифта к тексту перевода"""
        self.trans_text.configure(font=("Segoe UI", self.trans_font_size))

    # ════════════════════════════════════════════════════════════════════════
    # ПРАВЫЙ КЛИК
    # ════════════════════════════════════════════════════════════════════════

    def _on_right_click(self, event):
        if not self.selection_rect:
            return
        text = self._get_selection_text()
        if text.strip():
            self._show_floating_translation(event, text)

    def _show_floating_translation(self, event, text):
        self._close_popup()
        self.popup_win = tk.Toplevel(self.root)
        self.popup_win.title("Перевод")
        self.popup_win.overrideredirect(True)
        self.popup_win.attributes("-topmost", True)
        x, y = event.x_root + 15, event.y_root + 15
        self.popup_win.geometry(f"450x350+{x}+{y}")
        frame = tk.Frame(self.popup_win, bg="#f5f5f5", relief=tk.RAISED, bd=1)
        frame.pack(fill=tk.BOTH, expand=True)
        tk.Label(frame, text="Оригинал:", font=("Arial", 9, "bold"), bg="#f5f5f5").pack(anchor=tk.W, padx=8, pady=2)
        tk.Label(frame, text=text[:500], wraplength=420, justify=tk.LEFT, bg="#f5f5f5").pack(fill=tk.X, padx=8, pady=2)
        tk.Label(frame, text="Перевод:", font=("Arial", 9, "bold"), bg="#f5f5f5").pack(anchor=tk.W, padx=8, pady=2)
        trans_label = tk.Label(frame, text="⏳ Загрузка...", wraplength=420, justify=tk.LEFT, bg="#ffffcc")
        trans_label.pack(fill=tk.BOTH, expand=True, padx=8, pady=5)
        tk.Button(frame, text="✖", command=self._close_popup, bg="#ff6666", fg="white", width=2, relief=tk.FLAT).place(x=425, y=2)

        def do_translate():
            result = self._do_translate(text)
            trans_label.config(text=result, bg="#ffffcc")
        threading.Thread(target=do_translate, daemon=True).start()

        def close_on_click(e):
            if self.popup_win and not (self.popup_win.winfo_containing(e.x_root, e.y_root) == self.popup_win):
                self._close_popup()
        self.root.bind("<Button-1>", close_on_click, add="+")

    def _close_popup(self):
        if self.popup_win:
            try:
                self.popup_win.destroy()
            except:
                pass
            self.popup_win = None

    # ════════════════════════════════════════════════════════════════════════
    # КНОПКИ ПЕРЕВОДА ВЫДЕЛЕННОГО
    # ════════════════════════════════════════════════════════════════════════

    def translate_selection(self):
        text = self._get_selection_text()
        if not text:
            messagebox.showwarning("Внимание", "Выделите область на PDF")
            return
        self._open_result_window("Перевод выделенного", text, "⏳ Загрузка...")
        threading.Thread(target=self._do_translate_selection, args=(text,), daemon=True).start()

    def _do_translate_selection(self, text):
        result = self._do_translate(text)

        def update_ui():
            if hasattr(self, '_result_win_text') and self._result_win_text:
                self._result_win_text.config(state=tk.NORMAL)
                self._result_win_text.delete("1.0", tk.END)
                self._result_win_text.insert(1.0, result)
                self._result_win_text.config(state=tk.NORMAL)
        self.root.after(0, update_ui)

    def explain_ai(self):
        text = self._get_selection_text()
        if not text:
            messagebox.showwarning("Внимание", "Выделите область на PDF")
            return
        api_key = self.deepseek_key.get().strip()
        if not api_key:
            messagebox.showwarning("Внимание", "Введите DeepSeek API-ключ")
            return
        self._open_result_window("Объяснение от ИИ", text, "⏳ Отправляю запрос...")
        threading.Thread(target=self._do_explain_ai, args=(text, api_key), daemon=True).start()

    def _do_explain_ai(self, text, api_key):
        prompt = f"Перескажи этот текст своими словами на русском, кратко и понятно. Только пересказ, без лишних слов. Текст: {text}"
        try:
            resp = requests.post(DEEPSEEK_URL,
                                 headers={"Content-Type": "application/json", "Authorization": f"Bearer {api_key}"},
                                 json={"model": DEEPSEEK_MODEL, "max_tokens": 1024,
                                       "messages": [{"role": "user", "content": prompt}]},
                                 timeout=60)
            resp.raise_for_status()
            result = resp.json()["choices"][0]["message"]["content"].strip()
        except Exception as e:
            result = f"[Ошибка: {e}]"

        def update_ui():
            if hasattr(self, '_result_win_text') and self._result_win_text:
                self._result_win_text.config(state=tk.NORMAL)
                self._result_win_text.delete("1.0", tk.END)
                self._result_win_text.insert(1.0, f"Оригинал:\n{text}\n\n{'─'*50}\n\nОбъяснение ИИ:\n{result}")
                self._result_win_text.config(state=tk.NORMAL)
        self.root.after(0, update_ui)

    def _open_result_window(self, title, original_text, translated_text=""):
        """Окно с оригиналом и переводом"""
        win = tk.Toplevel(self.root)
        win.title(title)
        win.geometry("1000x600")
        win.configure(bg=self.colors["bg"])

        main_frame = tk.Frame(win, bg=self.colors["bg"])
        main_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

        paned = ttk.PanedWindow(main_frame, orient=tk.HORIZONTAL)
        paned.pack(fill=tk.BOTH, expand=True)

        # Левая панель — оригинал
        left_frame = ttk.Frame(paned)
        paned.add(left_frame, weight=1)
        ttk.Label(left_frame, text="📖 ОРИГИНАЛ", font=("Segoe UI", 11, "bold")).pack(anchor=tk.W, padx=5, pady=5)

        left_text = tk.Text(left_frame, wrap=tk.WORD, bg="#f5f5f5", fg="#111111",
                            font=("Segoe UI", 11), padx=10, pady=10, relief=tk.FLAT)
        left_text.pack(fill=tk.BOTH, expand=True)
        left_text.insert(1.0, original_text)
        left_text.config(state=tk.DISABLED)

        ttk.Button(left_frame, text="📋 Копировать оригинал",
                   command=lambda: self._copy_to_clipboard(original_text)).pack(pady=5)

        # Правая панель — перевод
        right_frame = ttk.Frame(paned)
        paned.add(right_frame, weight=1)
        ttk.Label(right_frame, text="🌐 ПЕРЕВОД", font=("Segoe UI", 11, "bold")).pack(anchor=tk.W, padx=5, pady=5)

        self._result_win_text = tk.Text(right_frame, wrap=tk.WORD, bg="#fffff0", fg="#111111",
                                        font=("Segoe UI", 11), padx=10, pady=10, relief=tk.FLAT)
        self._result_win_text.pack(fill=tk.BOTH, expand=True)
        self._result_win_text.insert(1.0, translated_text if translated_text else "⏳ Загрузка...")
        self._result_win_text.config(state=tk.NORMAL)

        btn_frame = ttk.Frame(right_frame)
        btn_frame.pack(pady=5)
        ttk.Button(btn_frame, text="📋 Копировать перевод",
                   command=lambda: self._copy_to_clipboard(self._result_win_text.get(1.0, tk.END).strip())).pack(side=tk.LEFT, padx=5)
        ttk.Button(btn_frame, text="📋 Копировать всё",
                   command=lambda: self._copy_to_clipboard(f"ОРИГИНАЛ:\n{original_text}\n\nПЕРЕВОД:\n{self._result_win_text.get(1.0, tk.END).strip()}")).pack(side=tk.LEFT, padx=5)

        win.grab_set()

    # ════════════════════════════════════════════════════════════════════════
    # СОХРАНЕНИЕ
    # ════════════════════════════════════════════════════════════════════════

    def save_docx(self):
        content = self.trans_text.get("1.0", tk.END).strip()
        if not content:
            messagebox.showwarning("Внимание", "Нет переведённого текста")
            return
        path = filedialog.asksaveasfilename(defaultextension=".docx", filetypes=[("Word документ", "*.docx")])
        if not path:
            return
        try:
            doc = python_docx.Document()
            doc.add_heading("Перевод PDF", level=1)
            for para in content.split("\n\n"):
                if para.strip():
                    doc.add_paragraph(para.strip())
            doc.save(path)
            self._set_status(f"Сохранено: {path}")
            messagebox.showinfo("Готово", f"Файл сохранён:\n{path}")
        except Exception as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить:\n{e}")

    # ════════════════════════════════════════════════════════════════════════
    # ПРОКРУТКА
    # ════════════════════════════════════════════════════════════════════════

    def _pdf_yscroll_set(self, first, last):
        if self.sync_scroll_var.get() and not self._scrolling:
            self._scrolling = True
            self.trans_text.yview_moveto(first)
            self._scrolling = False

    def _trans_yscroll_set(self, first, last):
        if self.sync_scroll_var.get() and not self._scrolling:
            self._scrolling = True
            self.pdf_canvas.yview_moveto(first)
            self._scrolling = False

    def _on_pdf_scroll_y(self, *args):
        self.pdf_canvas.yview(*args)

    def _on_trans_scroll_y(self, *args):
        self.trans_text.yview(*args)

    def _on_mousewheel_pdf(self, event):
        if event.num == 4:
            delta = -1
        elif event.num == 5:
            delta = 1
        else:
            delta = -1 if event.delta > 0 else 1
        if event.state & 0x4:
            self.zoom_in() if delta < 0 else self.zoom_out()
        else:
            self.pdf_canvas.yview_scroll(delta * 3, "units")
            if self.sync_scroll_var.get():
                self.trans_text.yview_scroll(delta * 3, "units")

    # ════════════════════════════════════════════════════════════════════════
    # ВСПОМОГАТЕЛЬНЫЕ
    # ════════════════════════════════════════════════════════════════════════

    def _set_status(self, msg):
        self.root.after(0, lambda: self.status_bar.config(text=msg))

    def _copy_to_clipboard(self, text):
        self.root.clipboard_clear()
        self.root.clipboard_append(text)
        self._set_status(f"Скопировано: {text[:50]}...")


# ════════════════════════════════════════════════════════════════════════════
if __name__ == "__main__":
    root = tk.Tk()
    try:
        root.tk.call("tk", "scaling", 1.3)
    except Exception:
        pass
    app = PDFTranslatorApp(root)
    root.mainloop()